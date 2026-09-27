#!/usr/bin/env python3
"""
Generate SWARMOS_Project_Blueprint_and_Feature_Inventory.docx

A self-contained project blueprint written specifically so that an EXTERNAL AI
reviewer can critique it without needing any other file. Regenerate this
document whenever the feature inventory or build status changes.

Usage:
    /pkg/OSS-python-/3.12.0/x86_64-linux4.18-glibc2.28/bin/python3 \
        docs/gen_project_blueprint_docx.py

Output:
    SWARMOS_Project_Blueprint_and_Feature_Inventory.docx   (repo root)
"""

from __future__ import annotations

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x59, 0x59, 0x59)
WARN = RGBColor(0x9C, 0x50, 0x00)


# --------------------------------------------------------------------------
# low level helpers
# --------------------------------------------------------------------------

def h1(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = ACCENT
    return p


def h2(doc, text):
    return doc.add_heading(text, level=2)


def h3(doc, text):
    return doc.add_heading(text, level=3)


def para(doc, text, size=10.5, italic=False, color=None, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    r.bold = bold
    if color is not None:
        r.font.color.rgb = color
    return p


def bullet(doc, text, level=0, size=10.5):
    style = "List Bullet" if level == 0 else "List Bullet %d" % (level + 1)
    try:
        p = doc.add_paragraph(style=style)
    except KeyError:
        p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def numbered(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def mono(doc, text, size=9):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(size)
    return p


def callout(doc, text, color=WARN):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.bold = True
    r.font.color.rgb = color
    return p


def table(doc, headers, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(htxt)
        r.bold = True
        r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(val))
            r.font.size = Pt(size)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


# --------------------------------------------------------------------------
# content data
# --------------------------------------------------------------------------

SOURCE_AUDIT = [
    ("SWARMOS_Product_Grade_UIUX_and_Simulation_Master_Directive.docx",
     "UI/UX authority + quality bar",
     "Principles only. Verified keyword counts over raw XML: warehouse=0, "
     "robot=0, aisle=0, battery=0, collision=0, deadlock=0, auction=0, "
     "tick=0, seed=0, WebSocket=0, endpoint=0. Zero hex colours, zero font "
     "names, zero px values. Embeds word/media/image1.jpg captioned "
     "'Negative reference supplied by the project owner. Do not reproduce "
     "this visual language.'",
     "Quality bar, prohibitions, product-grade mandate"),
    ("SWARMOS_Final_God_Mode_Master_Prompt_for_Claude.pdf",
     "Frontend feature inventory",
     "39 pages, ~90 sections. A prompt for Claude, not an architecture spec. "
     "Explicitly forbids touching backend/API layers. No endpoints, no WS "
     "topics, no DB tables, no file paths, no class names.",
     "Screen list, command palette, global search, focus mode, map layers, "
     "system health, MONITOR-UNDERSTAND-DECIDE-ACT-VERIFY backbone"),
    ("SWARMOS_Frontend_UIUX_Brief.pdf",
     "Frontend brief",
     "Screens, panels, theme and interaction intent. No contracts.",
     "Supporting UI requirements"),
    ("Medha_SWARMOS_Fresh_AI_Reconstruction_Specification.docx",
     "M2 Backend + M3 AI/ML",
     "The most substantive teammate document. Honestly marks lost items as "
     "UNDEFINED / NEEDS CONFIRMATION. States the original 'Step 27' 46-feature "
     "ML specification survives only as CATEGORY COUNTS, not feature names, "
     "and warns a later 46-feature list must NOT be treated as source of truth.",
     "Backend and ML requirements, explicit gap list"),
    ("Harshita_SWARMOS_Individual_Progress_Snapshot.docx",
     "M6 Database / simulation support",
     "8 short prose sections. No schemas, no code, no file paths, no table "
     "definitions.",
     "Role and rough status only"),
    ("Kushi_SWARMOS_Complete_Project_Handover_Prompt.docx",
     "(intended) M5 Simulation",
     "NOT a handover document. It is an 852-line meta-prompt instructing an AI "
     "on what a handover document should contain. Never states what the author "
     "actually built or owns.",
     "Nothing usable about the module"),
    ("SWARMOS_Final_God_Mode_Master_Prompt_for_Claude.zip",
     "-",
     "Pure duplicate. MD5-identical repack of the PDF plus the three docx "
     "files already present beside it.",
     "Nothing new"),
    ("app/coordination/ (existing code, Member 4)",
     "M4 Coordination",
     "The ONLY hard, executable, test-covered artefact in the project. 8 "
     "modules plus 6 pytest files. Commit 6ac7659 'feat: complete Step 4A "
     "decentralized auction coordination'.",
     "Canonical data model and message envelope for the whole system"),
]

TEAM = [
    ("M1", "Frontend / UI-UX", "Nandita",
     "Dashboard, warehouse visualization, robot status, charts"),
    ("M2", "Backend", "Medha",
     "FastAPI, REST/WS APIs, task and fleet management, connects all modules"),
    ("M3", "AI / ML", "Medha",
     "Dataset, model training, MOVE / WAIT / REROUTE / CHARGE predictions"),
    ("M4", "Robotics / Multi-Agent", "Adithya",
     "Robot-to-robot communication, collision avoidance, conflicts, deadlocks"),
    ("M5", "Simulation", "Kushi",
     "Virtual warehouse, 50 robots, robot movement and scenarios"),
    ("M6", "Database / Analytics", "Harshita",
     "Store robot/task data, compute performance, efficiency and reports"),
]

STACK = [
    ("Language / runtime", "Python", "3.12.0",
     "Verified present. NOTE: default `python3` on this host is 3.4.1 and must "
     "never be used."),
    ("Data contracts", "pydantic", "2.13.5", "Canonical models and validation"),
    ("Web framework", "FastAPI", "0.104.1", "REST + WebSocket + static serving"),
    ("ASGI server", "uvicorn", "0.52.4", "Single-process app host"),
    ("WebSocket", "websockets", "17.1", "Fleet telemetry stream"),
    ("Numerics", "numpy", "1.26.2", "Simulation and feature maths"),
    ("Dataframes", "pandas", "2.1.3", "Analytics and dataset assembly"),
    ("ML", "scikit-learn", "1.3.2", "Advisory action classifier"),
    ("Model persistence", "joblib", "1.3.2", "Pre-trained model shipped in repo"),
    ("Database", "sqlite3", "stdlib", "Zero-install persistence and traces"),
    ("Frontend", "Vanilla JS + Canvas 2D", "ES2022",
     "Deliberate choice: no npm, no build step, no node_modules at demo time"),
    ("Tests", "pytest", "9.1.1", "Per-module suites plus determinism test"),
    ("Docs", "python-pptx / python-docx", "1.0.2 / 1.1.0",
     "Script-generated deliverables"),
]

CONTRACTS = [
    ("Position", "x: float, y: float",
     "2-D warehouse coordinate in METRES. Both fields reject NaN and +/-Inf."),
    ("RobotStatus", "AVAILABLE | MOVING | WAITING | BLOCKED | FAILED | CHARGING",
     "str Enum. Controlled baseline states for an AMR."),
    ("MovementIntent",
     "target: Position, path: list[Position], eta: Optional[float] (>=0), "
     "intent_id: str (min_length=1), path_version: int (>=1)",
     "A robot's intended movement. path_version increments on every re-plan, "
     "enabling stale-intent detection without diffing waypoint lists."),
    ("AMRState",
     "robot_id: str, timestamp: float (unix epoch s), position: Position, "
     "velocity: float (>=0 m/s), heading: float ([0, 2*pi) rad), "
     "status: RobotStatus, battery: float (0-100), "
     "current_task_id: Optional[str], movement_intent: Optional[MovementIntent]",
     "Canonical per-robot state snapshot. Produced by M5, read by everyone."),
    ("MessageType", "ROBOT_STATE | PATH_INTENT (+ new types added in Step 4B)",
     "Controlled message type discriminator on the envelope."),
    ("CoordinationMessage",
     "schema_version: str (default '1.0'), message_id: str, type: MessageType, "
     "sender_id: str, timestamp: float, sequence: int (>=0), "
     "target_id: Optional[str] (None = broadcast), payload: dict",
     "Transport-agnostic envelope. message_id dedups replays, sequence detects "
     "out-of-order delivery, schema_version gates evolution."),
]

CONVENTIONS = [
    ("Coordinate system", "Continuous 2-D metres",
     "M5 owns any grid<->metre conversion. No other module converts."),
    ("Heading", "radians, [0, 2*pi), 0 = +X axis", "Validated in AMRState."),
    ("Velocity", "scalar speed in m/s, non-negative", "Not a vector."),
    ("Battery", "percent, 0.0 - 100.0 inclusive", "Validated."),
    ("Tick rate", "10 Hz (100 ms)",
     "DECIDED (authored). M2 owns the loop and is the only ticker."),
    ("Robot ID format", "R### (e.g. R001, R027, R050)", "DECIDED (authored)."),
    ("Task ID format", "T### (e.g. T001, T145)", "DECIDED (authored)."),
    ("Authoritative robot state", "M5 simulation",
     "Single writer. Everyone else reads."),
    ("Authoritative safety verdict", "M4 coordination",
     "Binding. Nothing may override it, including ML."),
    ("ML authority", "ADVISORY ONLY, never in the safety path",
     "A slow, absent or wrong prediction must degrade to rule-based behaviour."),
    ("Schema evolution", "envelope schema_version, starting '1.0'",
     "Additive changes only within a major version."),
]

# feature inventory: (id, feature, module, novelty, demo value, status)
FEATURES = [
    # M5 simulation
    ("S-01", "Warehouse floor model: aisles, racks, pick stations, charge docks",
     "M5", "No", "Foundation of the twin", "Planned"),
    ("S-02", "Robot kinematics in metres: speed, acceleration, turn rate, footprint",
     "M5", "No", "Believable motion", "Planned"),
    ("S-03", "Battery model with drain by motion and idle, plus charge cycles",
     "M5", "No", "Enables CHARGE decisions", "Planned"),
    ("S-04", "Scale to 50 concurrent robots", "M5", "No",
     "Meets the problem statement", "Planned"),
    ("S-05", "Seeded RNG so a seed reproduces a run exactly", "M5", "Yes (N3)",
     "Auditability - rare in swarm demos", "Planned"),
    ("S-06", "Named scenarios: baseline_10, rush_50, narrow_aisle_deadlock, "
     "robot_failure, battery_crisis, blocked_aisle", "M5", "Partly",
     "Each scenario proves one novelty", "Planned"),
    ("S-07", "Live fault injection: kill a robot, block an aisle, drain a battery",
     "M5", "Yes", "Judges love breaking things on stage", "Planned"),
    ("S-08", "Time controls: pause, step, 1x/2x/5x/10x speed", "M5", "No",
     "Demo pacing control", "Planned"),
    ("S-09", "Order/task generator with configurable arrival rate", "M5", "No",
     "Creates the workload", "Planned"),
    # M4 coordination
    ("C-01", "Decentralized auction task allocation - robots bid, no central "
     "dispatcher", "M4", "Yes (N1)",
     "Core differentiator vs central schedulers", "Step 4A DONE"),
    ("C-02", "Utility breakdown per bid - explainable, weighted components",
     "M4", "Yes", "Powers the Decision Inspector", "Step 4A DONE"),
    ("C-03", "Peer registry with heartbeat and liveness", "M4", "No",
     "Detects dead robots", "Step 4A DONE"),
    ("C-04", "Intent broadcast - robots publish MovementIntent before moving",
     "M4", "Yes (N2)", "Predictive, not reactive, collision avoidance",
     "Step 4A DONE"),
    ("C-05", "Space-time reservation of path cells", "M4", "Yes (N2)",
     "Conflicts resolved before motion", "Step 4A DONE"),
    ("C-06", "Conflict detection between competing intents", "M4", "Yes",
     "Visible conflict flashes on the map", "Step 4A DONE"),
    ("C-07", "Graded verdict ladder PROCEED -> WAIT -> YIELD -> SLOW -> REROUTE",
     "M4", "Yes (N4)", "Provable liveness policy, not ad-hoc backoff",
     "Planned (Step 4B)"),
    ("C-08", "Deadlock detection via wait-for cycle search", "M4", "Yes (N4)",
     "Resolves the narrow_aisle_deadlock scenario on camera", "Planned"),
    ("C-09", "Degraded-mode ladder: FULL -> NO_ML -> NO_NETWORK (local rules)",
     "M4", "Yes (N6)", "Edge resilience, BEL-relevant", "Planned"),
    ("C-10", "ML advisory intake with veto authority (safety firewall)",
     "M4", "Yes (N5)", "Answers 'can you trust the AI'", "Planned"),
    ("C-11", "Deterministic verdict ordering under identical seed", "M4",
     "Yes (N3)", "Same seed -> identical trace hash", "Planned"),
    # M3 ML
    ("A-01", "Dataset generated from real simulated runs, not synthetic noise",
     "M3", "Partly", "Defensible provenance", "Planned"),
    ("A-02", "Action classifier: MOVE / WAIT / REROUTE / CHARGE with confidence",
     "M3", "No", "The 'Edge-AI' in the problem title", "Planned"),
    ("A-03", "Hard inference timeout and confidence threshold", "M3", "Yes (N5)",
     "Makes the advisory boundary real, not rhetorical", "Planned"),
    ("A-04", "Pre-trained model committed so the demo never trains live",
     "M3", "No", "Demo reliability", "Planned"),
    ("A-05", "Agreement metric: how often ML agrees with the rule engine",
     "M3", "Yes", "Honest, measurable ML value claim", "Planned"),
    # M6 database and analytics
    ("D-01", "sqlite persistence: robots, tasks, task_events, conflicts, "
     "decisions, tick_metrics, runs", "M6", "No", "Evidence store", "Planned"),
    ("D-02", "KPIs: throughput, task latency, utilisation, idle %, conflict "
     "rate, deadlocks resolved, energy per task, distance", "M6", "No",
     "Turns a demo into evidence", "Planned"),
    ("D-03", "Baseline vs SWARMOS benchmark harness", "M6", "Yes",
     "The single most persuasive slide for judges", "Planned"),
    ("D-04", "Full run trace recording for replay", "M6", "Yes (N3)",
     "Demo insurance plus auditability", "Planned"),
    ("D-05", "Closed-loop analytics feeding KPIs back into bid weights",
     "M6", "Yes (N8)", "Optimization loop, not just a report", "Planned"),
    # M2 backend
    ("B-01", "10 Hz orchestration tick: sim -> coordination -> ML -> db -> fanout",
     "M2", "No", "The spine", "Planned"),
    ("B-02", "REST API: health, fleet, robots, tasks, scenarios, sim control, "
     "faults, metrics, report, benchmark, replay", "M2", "No",
     "Control surface", "Planned"),
    ("B-03", "WebSocket /ws/fleet broadcasting the fleet snapshot at 10 Hz",
     "M2", "No", "Live twin data", "Planned"),
    ("B-04", "Real health telemetry: tick duration, latency, data freshness, "
     "client count", "M2", "Yes", "Directive forbids faking health", "Planned"),
    ("B-05", "Graceful subsystem timeout handling (ML, db)", "M2", "Yes (N6)",
     "Feeds the degraded-mode ladder", "Planned"),
    # M1 frontend
    ("U-01", "Design-token layer: palette, type scale, 8px spacing grid, "
     "elevation, radius. No ad-hoc hex anywhere", "M1", "No",
     "The anti-'vibe-coded' foundation", "Planned"),
    ("U-02", "Live Map digital twin: racks, 50 robots with heading, intent "
     "trails, reservations, conflict flashes, verdict badges", "M1", "Yes (N7)",
     "The headline visual", "Planned"),
    ("U-03", "60 fps at 50 robots via layered canvases + rAF loop reading a "
     "mutable ref (no DOM churn)", "M1", "Yes",
     "Performance is itself a differentiator", "Planned"),
    ("U-04", "Map layer toggles: Robots, Routes, Trails, Tasks, Incidents, "
     "Congestion", "M1", "No", "Operator control", "Planned"),
    ("U-05", "Pan, zoom, fit-fleet, focus-robot", "M1", "No",
     "Navigation basics done properly", "Planned"),
    ("U-06", "Focus mode: selected robot prominent, context retained, "
     "unrelated elements quieter", "M1", "Yes", "Reduces cognitive load",
     "Planned"),
    ("U-07", "Command palette (Ctrl/Cmd+K) where every command performs a "
     "real action", "M1", "Yes", "Reads as a real product, not a demo",
     "Planned"),
    ("U-08", "Global search across robots, tasks, incidents, scenarios, "
     "locations with centre-and-select behaviour", "M1", "Yes",
     "Operator speed", "Planned"),
    ("U-09", "Fleet screen: sortable, filterable robot table plus detail "
     "inspector", "M1", "No", "Depth on demand", "Planned"),
    ("U-10", "Incidents screen with investigate -> act -> verify flow",
     "M1", "Yes", "Implements the MONITOR-UNDERSTAND-DECIDE-ACT-VERIFY "
     "backbone", "Planned"),
    ("U-11", "Decision Inspector: WHY this verdict, with the real utility "
     "breakdown numbers", "M1", "Yes (N1/N4)",
     "Explainability judges can interrogate", "Planned"),
    ("U-12", "Simulation Lab: scenario, robot count, seed, faults, speed",
     "M1", "Yes", "Lets judges drive the system themselves", "Planned"),
    ("U-13", "Analytics screen: purposeful charts only, all from M6",
     "M1", "No", "Evidence", "Planned"),
    ("U-14", "Benchmarks screen: baseline vs SWARMOS side by side", "M1",
     "Yes (N8)", "The money shot", "Planned"),
    ("U-15", "Replay screen with timeline scrubber over a recorded trace",
     "M1", "Yes (N3)", "Demo insurance and auditability", "Planned"),
    ("U-16", "System health panel with only real values, never faked",
     "M1", "Yes", "Directive mandate", "Planned"),
    ("U-17", "Designed empty, loading, error and degraded states for every "
     "panel", "M1", "Yes", "The clearest amateur/product dividing line",
     "Planned"),
    ("U-18", "Keyboard-first: palette, '/' search, arrow robot cycling, "
     "space pause, visible focus rings, ARIA labels", "M1", "No",
     "Accessibility and polish", "Planned"),
    ("U-19", "Tabular numerals for all metrics so digits do not jitter while "
     "ticking", "M1", "Yes", "A detail judges feel without naming", "Planned"),
    ("U-20", "Purposeful motion only: 120-200 ms state transitions, zero "
     "decorative animation", "M1", "No", "Industrial HMI discipline",
     "Planned"),
]

NOVEL = [
    ("N1", "Decentralized auction task allocation",
     "Robots bid for tasks and the winner is decided by the peer group. There "
     "is no central dispatcher to fail.",
     "Most warehouse fleets use a central scheduler, which is a single point "
     "of failure and a scaling bottleneck.",
     "Kill the 'coordinator' mid-demo. The fleet keeps allocating tasks.",
     "M4 (owner); M2, M5 support"),
    ("N2", "Intent broadcast plus space-time reservation",
     "Every robot publishes its MovementIntent with a path_version before it "
     "moves. Conflicts are detected and resolved against future reservations, "
     "not against current proximity.",
     "Conventional AMRs avoid collisions reactively when sensors see something "
     "close. This resolves the conflict before motion begins.",
     "Show two robots approaching a junction; the conflict resolves before "
     "either enters it.",
     "M4 (owner); M5 support"),
    ("N3", "Deterministic replay",
     "The same seed reproduces a bit-identical run. Traces are recorded and "
     "can be scrubbed on a timeline.",
     "Multi-agent demos are usually irreproducible, which makes them "
     "unauditable and undebuggable.",
     "Run seed 42 twice, show identical trace hashes; then scrub the replay.",
     "M4 + M5 (co-owners); M6 stores traces"),
    ("N4", "Deadlock detection with a graded escalation ladder",
     "PROCEED -> WAIT -> YIELD -> SLOW -> REROUTE, with wait-for cycle "
     "detection to break true deadlocks.",
     "Ad-hoc random backoff is the usual approach and it cannot guarantee "
     "liveness. This is an explicit, inspectable policy.",
     "Run narrow_aisle_deadlock and watch the ladder escalate and resolve it.",
     "M4 (owner)"),
    ("N5", "Edge-AI advisory layer behind a hard safety firewall",
     "The ML model suggests MOVE / WAIT / REROUTE / CHARGE with a confidence "
     "score. The coordination engine can veto it. Timeouts and low confidence "
     "fall back to rules automatically.",
     "Teams typically either put ML in the control path (unsafe) or bolt it on "
     "cosmetically (pointless). This gets the benefit without the risk, and "
     "the boundary is demonstrable.",
     "Disable the model live. The fleet continues correctly at slightly lower "
     "efficiency, and the UI shows mode NO_ML.",
     "M3 (owner); M4 enforces the firewall"),
    ("N6", "Graceful degradation ladder",
     "FULL -> NO_ML -> NO_NETWORK. In the worst case each robot runs local "
     "rules only and the fleet still operates safely.",
     "Directly addresses edge-deployment reality, which matters to a defence "
     "PSU sponsor like BEL.",
     "Toggle each mode on stage and show the fleet degrade rather than stop.",
     "M4 (owner); M2, M5 support"),
    ("N7", "Live digital-twin command centre",
     "A product-grade operational UI that renders not just robot positions but "
     "the coordination reasoning: intents, reservations, conflicts, verdicts.",
     "Competing demos show moving dots. This shows WHY each dot moved.",
     "Open the Decision Inspector on any robot mid-run and read the actual "
     "utility numbers.",
     "M1 (owner); M4, M5 supply data"),
    ("N8", "Closed-loop analytics",
     "KPIs computed by M6 feed back into the auction bid weights, so the fleet "
     "tunes itself against measured performance.",
     "Analytics are normally a read-only report. Here they close a control "
     "loop.",
     "Show the benchmark screen: baseline vs SWARMOS with the loop enabled and "
     "disabled.",
     "M6 (owner); M4, M3 support"),
]

VALUE = [
    ("Nandita", "M1 Frontend / UI-UX", "N7",
     "N4 (visualises the ladder), N8 (charts), N3 (replay UI)",
     "Makes invisible coordination legible. Without this the novelty is "
     "invisible to judges."),
    ("Medha", "M2 Backend", "(enabler for all 8)",
     "N5 (serves inference), N6 (timeout and fallback), N1 (task lifecycle)",
     "The integration spine. Without M2 there are modules but no system."),
    ("Medha", "M3 AI / ML", "N5",
     "N1 (bid-cost features), N8 (feeds the loop)",
     "Supplies the 'Edge-AI' the problem statement asks for."),
    ("Adithya", "M4 Robotics / Multi-Agent", "N1, N2, N3, N4, N6",
     "N5 (enforces the firewall), N7 (supplies overlay data)",
     "The technical core and the safety guarantee. Most of the novelty."),
    ("Kushi", "M5 Simulation", "50-robot scenario suite; co-owns N3",
     "N1, N2, N6 (generates the stress cases)",
     "The only way to prove any claim at scale."),
    ("Harshita", "M6 Database / Analytics", "N8",
     "N3 (persists traces), N7 (chart data)",
     "Turns a demo into measured evidence."),
]

ANTI_VIBE = [
    "A real design-token layer in web/styles/tokens.css. Every colour, size, "
    "space, radius and elevation in the UI references a token. No ad-hoc hex "
    "values anywhere in the codebase.",
    "No emoji in the UI. No gradient soup. No purple-on-black startup "
    "landing-page aesthetic. The target is dense, calm, high-contrast "
    "industrial HMI.",
    "Typography discipline: one UI family plus one tabular-numeric mono for "
    "all metrics, and a strict six-step type scale. Metrics must not jitter "
    "as digits change.",
    "NO FAKE DATA, EVER. Every number on screen traces to a real value "
    "produced by the engine. No placeholder sparklines, no lorem text, no "
    "hardcoded percentages. If a value is unavailable the UI shows an explicit "
    "empty or degraded state.",
    "Purposeful motion only. 120-200 ms easing on state transitions. Zero "
    "decorative animation. Robots move because the simulation moved them.",
    "Every control performs a real action. No dead buttons, no toggles that "
    "only change their own colour, no commands that log to console.",
    "Designed empty, loading, error and degraded states for every panel - not "
    "just the happy path. This is the clearest dividing line between a product "
    "and a demo.",
    "Keyboard-first operation: Ctrl/Cmd+K palette, '/' for search, arrow keys "
    "to cycle robots, space to pause, visible focus rings, ARIA labels.",
    "60 fps with 50 robots. Layered canvases separate the static floor from "
    "dynamic robots and overlays. The render loop reads a mutable state ref; "
    "the DOM is only touched for low-frequency aggregates.",
    "Honest system health: real WebSocket latency, real data freshness, real "
    "tick duration. Never faked, never simulated, never optimistic.",
]

REST_API = [
    ("GET", "/api/health", "Real liveness plus tick duration and client count"),
    ("GET", "/api/fleet", "Current fleet snapshot (all AMRState)"),
    ("GET", "/api/robots", "Robot list with status and task"),
    ("GET", "/api/robots/{robot_id}", "Single robot detail plus intent and verdict"),
    ("GET", "/api/tasks", "Task list with lifecycle state"),
    ("POST", "/api/tasks", "Create a task (pick station -> destination)"),
    ("GET", "/api/scenarios", "Available named scenarios"),
    ("POST", "/api/sim/start", "Start a scenario with robot count and seed"),
    ("POST", "/api/sim/stop", "Stop the run"),
    ("POST", "/api/sim/reset", "Reset to a clean state"),
    ("POST", "/api/sim/speed", "Set time multiplier (1x/2x/5x/10x) or pause"),
    ("POST", "/api/faults", "Inject a fault: kill robot, block aisle, drain battery"),
    ("POST", "/api/mode", "Force degraded mode: FULL | NO_ML | NO_NETWORK"),
    ("GET", "/api/metrics", "Live KPI set from M6"),
    ("GET", "/api/incidents", "Conflicts, deadlocks, failures"),
    ("GET", "/api/decisions/{robot_id}", "Verdict history plus utility breakdown"),
    ("GET", "/api/report", "Run report"),
    ("POST", "/api/benchmark", "Run baseline vs SWARMOS comparison"),
    ("GET", "/api/runs", "Recorded runs available for replay"),
    ("GET", "/api/replay/{run_id}", "Trace frames for the replay scrubber"),
    ("WS", "/ws/fleet", "10 Hz fleet snapshot: robots, intents, conflicts, "
     "verdicts, KPIs, mode, health"),
]

NFR = [
    ("Tick rate", "10 Hz, 100 ms budget per tick",
     "Measured and displayed; overruns are reported, not hidden"),
    ("Robot scale", "50 concurrent robots", "Problem-statement requirement"),
    ("Render", "60 fps at 50 robots", "Layered canvas, rAF, no DOM churn"),
    ("ML inference", "hard timeout, advisory only",
     "Exceeding budget must degrade to rules, never stall the tick"),
    ("Determinism", "identical seed -> identical trace hash",
     "Enforced by an automated test, not by inspection"),
    ("Safety invariant", "no two robots occupy overlapping space-time cells",
     "Asserted continuously in tests; violations fail the build"),
    ("Liveness invariant", "no robot waits indefinitely",
     "The escalation ladder must terminate; deadlock cycles are broken"),
    ("Degradation", "FULL / NO_ML / NO_NETWORK all operate safely",
     "Each mode is separately tested"),
    ("Install", "zero network install; no npm at demo time",
     "All dependencies verified present in the environment"),
]

GAPS = [
    "The original 46-feature ML specification is LOST. Only category counts "
    "survive. I will implement a documented, honest feature set and state "
    "clearly that it is a reconstruction, not the original.",
    "No teammate document defines the warehouse geometry, so I am authoring "
    "it. If a real BEL warehouse layout exists it should replace mine.",
    "Medha carries both M2 and M3, which is the schedule risk. Mitigation: M3 "
    "dataset generation is driven off M5 simulator logs so it does not block "
    "M2 API work.",
    "Kushi's document never states what she built, so M5 is being written from "
    "scratch rather than integrated.",
    "The ML value claim is the weakest novelty. If the model only agrees with "
    "the rule engine, the honest metric (A-05) will show that. I would rather "
    "report a small real gain than a large fake one.",
    "No hardware and no ROS 2. Everything is simulated. A judge may ask about "
    "the sim-to-real gap and we need a credible answer.",
    "Benchmarking against a 'baseline' that I also wrote is inherently "
    "self-serving. The baseline must be a genuinely reasonable centralised "
    "greedy allocator, not a strawman.",
]


# --------------------------------------------------------------------------
# document build
# --------------------------------------------------------------------------

def build(out_path: str) -> str:
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    # ---------------- title
    t = doc.add_heading("SWARMOS", level=0)
    for r in t.runs:
        r.font.color.rgb = ACCENT
    para(doc, "Project Blueprint and Complete Feature Inventory", size=15,
         bold=True)
    para(doc, "Edge-AI Based Distributed Fleet Coordination for Autonomous "
              "Mobile Robots (AMRs) in Smart Warehouses", size=12)
    para(doc, "Smart India Hackathon 2026  |  Problem ID SIH26123  |  "
              "Bharat Electronics Limited (BEL)  |  Theme: Robotics & Drones",
         size=10, color=MUTED)
    para(doc, "Document purpose: this is a self-contained blueprint prepared "
              "for INDEPENDENT AI REVIEW. It describes the entire intended "
              "system, every feature to be built, and every novelty claim, so "
              "that a reviewer can identify missing features, weak novelty "
              "claims and demo risks WITHOUT access to any other file or to "
              "the source code.", size=10, italic=True)
    doc.add_page_break()

    # ---------------- 1 context
    h1(doc, "1. Context and Objective")
    para(doc, "SWARMOS is a software system that coordinates a fleet of "
              "autonomous mobile robots (AMRs) inside a warehouse WITHOUT a "
              "central dispatcher. Robots negotiate task allocation among "
              "themselves, publish their movement intentions, reserve space "
              "and time along their paths, and resolve conflicts and deadlocks "
              "through an explicit escalation policy. An edge-AI model advises "
              "the fleet but is never allowed into the safety path.")
    h2(doc, "1.1 What winning requires")
    for x in [
        "A live demo that visibly works at 50 robots without stutter or crash.",
        "Novelty a domain expert recognises as genuinely non-standard, not "
        "merely a competent implementation of known techniques.",
        "Measured evidence: numbers comparing SWARMOS against a credible "
        "baseline, not adjectives.",
        "Explainability: when a judge asks WHY a robot waited, the UI answers "
        "with real computed values on screen.",
        "Robustness under deliberate abuse: the demo must survive a judge "
        "killing a robot, blocking an aisle or disabling the model.",
        "A product-grade interface. The system must not look like a student "
        "project or a generated template.",
    ]:
        bullet(doc, x)

    # ---------------- 2 source audit
    h1(doc, "2. Source-Material Audit (stated honestly)")
    para(doc, "Seven documents were supplied by teammates plus one existing "
              "codebase. Only two carry hard engineering content. This matters "
              "for the reviewer: most requirements below are AUTHORED by me, "
              "not inherited, so they are open to challenge.")
    table(doc,
          ["Source", "Intended scope", "What it actually contains",
           "Usable output"],
          SOURCE_AUDIT, size=8)
    callout(doc, "Consequence: the project is specified CONTRACT-FIRST from "
                 "the only executable artefact that exists (the M4 "
                 "coordination code). Everything else is authored to fit it.")

    # ---------------- 3 team
    h1(doc, "3. Team Structure and Ownership")
    table(doc, ["#", "Module", "Owner", "Scope"], TEAM, size=9)
    para(doc, "Six modules across five people. Medha owns both M2 and M3. For "
              "the current build phase all modules are being implemented "
              "centrally so that a complete, integrated system exists; "
              "ownership above remains the presentation and maintenance split.",
         italic=True)

    # ---------------- 4 architecture
    h1(doc, "4. System Architecture")
    h2(doc, "4.1 Tier model")
    mono(doc, """
+--------------------------------------------------------------+
 TIER 5  PRESENTATION      M1 FRONTEND (Vanilla JS + Canvas 2D)
                           digital twin | fleet | incidents |
                           sim lab | analytics | benchmarks |
                           replay | palette | search
+--------------------------^-----------------------------------+
                           |  REST (control) + WebSocket (10 Hz telemetry)
+--------------------------+-----------------------------------+
 TIER 4  ORCHESTRATION     M2 BACKEND (FastAPI + uvicorn)
                           owns THE tick loop | REST | /ws/fleet
+---^--------^---------^---------^-----------------------------+
    |        |         |         |
    | advisory|state   | persist | binding verdict
    | (vetoable)       |         |
+---+----+ +-+------+ +-+-----+ +-+----------------------------+
 M3 ML    | M5 SIM   | M6 DB   | M4 COORDINATION
 sklearn  | 50 AMRs  | sqlite  | auction | intent | reservation
 advisory | seeded   | KPIs    | conflict | deadlock | ladder
+---------+ +--------+ +-------+ +----------------------------+
+--------------------------------------------------------------+
 TIER 2  CONTRACTS   app/coordination/models.py + messages.py
                     schema_version "1.0"  (single source of truth)
+--------------------------------------------------------------+
 TIER 1  DATA PLANE  sqlite (durable) + in-memory tick state
+--------------------------------------------------------------+
""")
    h2(doc, "4.2 The four architectural laws")
    for n, x in enumerate([
        "M5 (simulation) is the SINGLE AUTHORITATIVE SOURCE of robot state. "
        "One writer; every other module reads.",
        "M4 (coordination) is the BINDING SAFETY ARBITER. Its verdict cannot "
        "be overridden by any other module, including the ML model.",
        "M3 (ML) is ADVISORY ONLY and is NEVER in the safety path. Absence, "
        "latency or error must degrade the system, never break it.",
        "ONE CANONICAL MODEL PER CONCEPT. app/coordination/models.py is the "
        "single definition of AMRState. No module redefines it.",
    ], 1):
        numbered(doc, x)

    h2(doc, "4.3 Technology stack (all versions verified present)")
    table(doc, ["Layer", "Technology", "Version", "Note"], STACK, size=9)
    callout(doc, "Deliberate decision: the frontend is vanilla JS + Canvas 2D, "
                 "NOT React/Vite. Rationale: no npm install and no build step "
                 "at demo time, no node_modules drift, and Canvas 2D handles "
                 "50 robots at 60 fps where a React re-render cycle would be "
                 "a liability. Reviewer: challenge this if you disagree.")

    # ---------------- 5 contracts
    h1(doc, "5. Canonical Data Contracts")
    para(doc, "These are the real, implemented pydantic models from the "
              "existing M4 codebase. They are the shared language of the "
              "entire system.")
    table(doc, ["Model", "Fields and constraints", "Meaning"], CONTRACTS,
          size=8)
    h2(doc, "5.1 Decided conventions")
    para(doc, "Where no source document specified a value, I decided one. Each "
              "is open to objection but none is left undefined, so "
              "implementation is never blocked.", italic=True)
    table(doc, ["Item", "Decision", "Rationale / authority"], CONVENTIONS,
          size=9)

    # ---------------- 6 feature inventory
    h1(doc, "6. Complete Feature Inventory")
    callout(doc, "REVIEWER: this is the critical section. Please diff this "
                 "list against what a winning SIH entry would contain and "
                 "tell me what is MISSING.", color=ACCENT)
    para(doc, "Status values: 'Step 4A DONE' means implemented and covered by "
              "passing tests today. 'Planned' means specified and scheduled in "
              "this build.")
    table(doc,
          ["ID", "Feature", "Module", "Novel?", "Demo value", "Status"],
          FEATURES, size=8)
    para(doc, "Total features enumerated: %d. Implemented today: %d. "
              "Planned in this build: %d."
              % (len(FEATURES),
                 sum(1 for f in FEATURES if "DONE" in f[5]),
                 sum(1 for f in FEATURES if "Planned" in f[5])),
         italic=True)

    # ---------------- 7 novelty
    h1(doc, "7. Targeted Novel Features")
    para(doc, "Eight novelty claims. Each is stated with its justification and "
              "the exact way it is demonstrated live, so a reviewer can judge "
              "whether the claim is defensible.")
    for nid, name, what, why, demo, owner in NOVEL:
        h2(doc, "%s  %s" % (nid, name))
        para(doc, "What it is: " + what)
        para(doc, "Why it is novel: " + why)
        para(doc, "How it is demonstrated: " + demo)
        para(doc, "Owner: " + owner, size=9, color=MUTED)
    callout(doc, "REVIEWER: please rank these by real judge impact and tell me "
                 "which are actually TABLE STAKES rather than novel. Be "
                 "brutal. I would rather cut a weak claim than have a judge "
                 "dismantle it.", color=ACCENT)

    # ---------------- 8 member value
    h1(doc, "8. How Each Member Adds Value")
    table(doc,
          ["Member", "Module", "Owns novelty", "Supports", "Value contributed"],
          VALUE, size=8)

    # ---------------- 9 UI/UX
    h1(doc, "9. UI/UX Specification")
    para(doc, "The supplied Master Directive is a principles document: it "
              "contains no colours, fonts or dimensions, and embeds a "
              "screenshot captioned 'Negative reference supplied by the "
              "project owner. Do not reproduce this visual language.' Its "
              "closing line is 'Build the product. Make the architecture "
              "visible. Make the Simulation unforgettable.' The design system "
              "below is therefore authored by me to meet that bar.")
    h2(doc, "9.1 Anti-amateur engineering rules")
    for x in ANTI_VIBE:
        bullet(doc, x)
    h2(doc, "9.2 Screens")
    for x in [
        "Live Map - the digital twin. Racks, aisles, 50 robots with heading, "
        "intent trails, reservation cells, conflict flashes, verdict badges, "
        "congestion heatmap.",
        "Fleet - sortable and filterable robot table with a detail inspector.",
        "Incidents - conflicts, deadlocks and failures with an explicit "
        "investigate -> act -> verify flow.",
        "Simulation Lab - scenario picker, robot count, seed entry, fault "
        "injection, speed multipliers, pause and step.",
        "Analytics - purposeful charts only, every series sourced from M6.",
        "Benchmarks - baseline versus SWARMOS side by side.",
        "Replay - timeline scrubber over a recorded trace.",
        "Decision Inspector (contextual) - why this robot received this "
        "verdict, with the real utility breakdown.",
        "Global layers - command palette, global search, focus mode, system "
        "health.",
    ]:
        bullet(doc, x)
    h2(doc, "9.3 Workflow backbone")
    para(doc, "Every major workflow follows MONITOR -> UNDERSTAND -> DECIDE -> "
              "ACT -> VERIFY. Example: an incident appears, the operator "
              "investigates it, understands the cause, acts, observes the "
              "result and verifies recovery.")
    h2(doc, "9.4 Questions the UI must answer at any moment")
    para(doc, "WHAT is happening?  WHERE?  WHICH robot?  WHICH task?  HOW "
              "serious?  WHY?  WHAT is SWARMOS doing?  WHAT can I do?  WHAT "
              "happened afterward?")

    # ---------------- 10 simulation
    h1(doc, "10. Simulation Specification")
    for x in [
        "Warehouse: rack blocks separated by aisles, pick stations along one "
        "edge, charge docks along another, with support for dynamically "
        "blocked cells. Continuous metre coordinates.",
        "Robot: differential-drive style kinematics with speed, acceleration "
        "and turn-rate limits, a circular footprint, and battery drain that "
        "differs between motion and idle.",
        "Scale: 50 concurrent robots, configurable down to 3 for the vertical "
        "slice.",
        "Scenarios: baseline_10, rush_50, narrow_aisle_deadlock, "
        "robot_failure, battery_crisis, blocked_aisle.",
        "Fault injection at runtime: kill a robot, block an aisle, drain a "
        "battery, disable the ML model, sever the network.",
        "Determinism: every stochastic decision draws from a single seeded "
        "RNG. The same seed reproduces a bit-identical run, verified by "
        "comparing trace hashes in an automated test.",
        "Time: 10 Hz simulation tick with pause, single-step and 1x/2x/5x/10x "
        "multipliers.",
        "Tasks: generated at a configurable arrival rate from pick station to "
        "destination, with a full lifecycle recorded by M6.",
    ]:
        bullet(doc, x)

    # ---------------- 11 ML
    h1(doc, "11. Machine Learning Specification")
    for x in [
        "Purpose: advise the coordination engine on the best next action for a "
        "robot - MOVE, WAIT, REROUTE or CHARGE - with a confidence score.",
        "Model: scikit-learn RandomForestClassifier, shipped pre-trained via "
        "joblib so the demo never trains live.",
        "Training data: generated by running the named scenarios through the "
        "rule-based coordination engine and recording (state -> chosen action) "
        "pairs. Provenance is real simulated behaviour, not synthetic noise.",
        "Features: derived from robot state, local congestion, task context, "
        "battery and conflict history. NOTE: the original 46-feature "
        "specification is LOST - only category counts survive - so this "
        "feature set is an honest reconstruction and is documented as such.",
        "Safety firewall: a hard inference timeout plus a confidence "
        "threshold. On timeout, exception or low confidence the predictor "
        "returns None and the coordination engine silently uses its rules.",
        "Honest metric: the system reports how often the model AGREES with the "
        "rule engine and what measurable improvement it produces. If the gain "
        "is small, the number shown will be small.",
    ]:
        bullet(doc, x)

    # ---------------- 12 analytics
    h1(doc, "12. Analytics, KPIs and Benchmarking")
    for x in [
        "Throughput - completed tasks per minute.",
        "Task latency - assignment to completion, mean and p95.",
        "Robot utilisation and idle percentage.",
        "Conflict rate - conflicts detected per robot-minute.",
        "Deadlocks detected and deadlocks resolved.",
        "Energy per completed task, and total distance travelled.",
        "Verdict distribution - how often each of PROCEED / WAIT / YIELD / "
        "SLOW / REROUTE was issued.",
        "Benchmark methodology: identical scenario, identical seed, identical "
        "task stream, run twice - once with a centralised greedy allocator as "
        "the baseline and once with SWARMOS. The baseline must be a genuinely "
        "reasonable implementation, not a strawman, or the comparison is "
        "worthless.",
    ]:
        bullet(doc, x)

    # ---------------- 13 API
    h1(doc, "13. Interface Contract")
    table(doc, ["Method", "Path", "Purpose"], REST_API, size=9)

    # ---------------- 14 NFR
    h1(doc, "14. Non-Functional Requirements")
    table(doc, ["Requirement", "Target", "How it is enforced"], NFR, size=9)

    # ---------------- 15 verification
    h1(doc, "15. Test and Verification Plan")
    for x in [
        "pytest suites per module: sim, db, coordination, ml, api.",
        "Determinism test: two runs at the same seed must produce identical "
        "trace hashes.",
        "Safety invariant test: no two robots may occupy overlapping "
        "space-time reservations at any tick.",
        "Liveness test: no robot waits beyond a bounded number of ticks; the "
        "escalation ladder must terminate.",
        "Degraded-mode tests: FULL, NO_ML and NO_NETWORK each exercised "
        "independently.",
        "API exercise: every REST endpoint called and every response schema "
        "validated; WebSocket connected and tick rate measured.",
        "Browser walkthrough: the running site loaded in a real browser and "
        "every screen, control, palette command and search path exercised by "
        "hand, with results reported honestly including anything broken.",
        "Performance measurement: actual frame rate recorded at 50 robots, "
        "reported as measured rather than claimed.",
    ]:
        bullet(doc, x)

    # ---------------- 16 build sequence
    h1(doc, "16. Build Sequence")
    for x in [
        "M5 simulation - warehouse, robots, scenarios, seeded determinism, "
        "time controls.",
        "M6 database and analytics - persistence, KPIs, reports, trace "
        "recording.",
        "M4 coordination extension (Step 4B) - verdict ladder, deadlock "
        "detection, degraded modes, ML intake with veto.",
        "M2 backend - the 10 Hz tick loop, REST surface, /ws/fleet.",
        "M1 frontend - design tokens, app shell, live map, then each screen, "
        "then palette, search and focus mode.",
        "M3 ML - dataset from real runs, training, advisory predictor behind "
        "the firewall.",
        "Benchmark harness, trace replay, polish.",
        "Full verification pass, then the master specification deck.",
    ]:
        numbered(doc, x)
    para(doc, "The system is kept runnable at the end of every step rather "
              "than integrated once at the end.", italic=True)

    # ---------------- 17 gaps
    h1(doc, "17. Known Gaps, Risks and Weak Points")
    para(doc, "Stated openly so that a reviewer attacks the real weaknesses "
              "instead of rediscovering them.", italic=True)
    for x in GAPS:
        bullet(doc, x)

    # ---------------- 18 request
    h1(doc, "18. Request to the Reviewer")
    for x in [
        "What high-impact NOVEL feature is missing from section 6 that a "
        "winning team would have?",
        "Which of the eight novelty claims in section 7 are actually table "
        "stakes? Rank them by real judge impact.",
        "Is anything in the SIH26123 problem statement not addressed?",
        "What single addition would most raise demo impact per unit of effort?",
        "Which questions would a hostile expert judge use to expose us, and "
        "how should we pre-empt them?",
        "What in section 9 would still read as amateur, and what is the "
        "specific fix?",
        "What should be CUT because it costs more effort than it earns in "
        "marks?",
        "Every competing team will show robots moving on a warehouse map. What "
        "makes this unmistakably better?",
    ]:
        numbered(doc, x)

    doc.add_paragraph()
    para(doc, "End of blueprint. Generated from docs/gen_project_blueprint_docx.py "
              "- regenerate after any change to the feature inventory.",
         size=9, italic=True, color=MUTED)

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    out = os.path.join(
        os.path.dirname(here),
        "SWARMOS_Project_Blueprint_and_Feature_Inventory.docx")
    path = build(out)
    d = Document(path)
    print("Wrote: %s" % path)
    print("Paragraphs: %d   Tables: %d" % (len(d.paragraphs), len(d.tables)))

# File contains AI-generated response based on internal company sources
