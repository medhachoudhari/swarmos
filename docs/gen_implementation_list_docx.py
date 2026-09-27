#!/usr/bin/env python3
"""
Generate SWARMOS_Final_Implementation_Feature_List.docx

The frozen build contract. Every feature that will be implemented, every feature
that was deliberately cut, every claim that will be made, and the acceptance
test that proves each one. Produced after four independent external AI
adversarial reviews were ingested and triaged.

Usage:
    /pkg/OSS-python-/3.12.0/x86_64-linux4.18-glibc2.28/bin/python3 \
        docs/gen_implementation_list_docx.py

Output:
    SWARMOS_Final_Implementation_Feature_List.docx   (repo root)
"""

from __future__ import annotations

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor, Inches

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x59, 0x59, 0x59)
WARN = RGBColor(0x9C, 0x50, 0x00)
GOOD = RGBColor(0x1E, 0x6B, 0x3A)


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------

def h1(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = ACCENT
    return p


def h2(doc, text):
    return doc.add_heading(text, level=2)


def para(doc, text, size=10.5, italic=False, color=None, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    r.bold = bold
    if color is not None:
        r.font.color.rgb = color
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def numbered(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def mono(doc, text, size=9):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(size)
    return p


def callout(doc, text, color=WARN):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.bold = True
    r.font.color.rgb = color
    return p


def table(doc, headers, rows, widths=None, size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(htxt)
        r.bold = True
        r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(size)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


# --------------------------------------------------------------------------
# DATA - the frozen scope
# --------------------------------------------------------------------------

# Already implemented and unit tested (M4 Step 4A).
DONE = [
    ("C-01", "Canonical data contracts: Position, RobotStatus, MovementIntent, "
             "AMRState with full validation", "M4",
     "tests/test_coordination_contracts.py"),
    ("C-02", "CoordinationMessage envelope with schema_version, message_id, "
             "sequence, target_id", "M4", "tests/test_coordination_contracts.py"),
    ("C-03", "Peer registry with liveness tracking", "M4",
     "tests/test_peer_registry.py"),
    ("C-04", "Intent registry keyed by robot with path_version invalidation",
     "M4", "tests/test_intent_registry.py"),
    ("C-05", "Space-time reservation table", "M4", "tests/test_reservation.py"),
    ("C-06", "Conflict detection over reservations; decentralised auction "
             "bidding", "M4", "tests/test_conflict.py, tests/test_auction.py"),
]

# The accepted build list. Wave 1 = S effort. Wave 2 = M effort demo carriers.
# Wave 3 = the mechanism items contributed by reviewers 3 and 4.
BUILD = [
    # --- Wave 1: S-effort, every one turns a claim into a number on screen
    ("X-09", "Safety invariant monitor", "M4 + M1",
     "Live counter of space-time overlap violations, held at 0, alongside total "
     "tick-cells checked. Always on screen in the status strip.",
     "Assert 0 violations across all scenarios at 10/25/50/100 robots",
     "S", 9, 1),
    ("X-04", "Edge compute budget panel", "M3 + M1",
     "Measured on-agent: model size in KB, inference latency p50/p95/p99, "
     "feature-extraction cost, ticks/s headroom against the 100 ms budget.",
     "Panel shows real measured numbers, not constants",
     "S", 9, 1),
    ("X-02", "Message-bus impairment injection", "M4",
     "Per-link drop probability, latency jitter distribution, and whole-zone "
     "partition. Live toggles.",
     "NO_NETWORK is entered because packets were actually dropped",
     "S", 9, 1),
    ("X-13", "Statistical benchmark rigour", "M6",
     "At least 10 seeds per configuration, mean with 95% confidence interval, "
     "paired per-seed deltas against the baseline.",
     "Benchmark screen shows CI error bars and paired deltas",
     "S", 8, 1),
    ("X-03", "Message budget accounting", "M4 + M6",
     "Messages per robot per tick, bytes/s, p95 consensus latency, all plotted "
     "against fleet size.",
     "Curve is flat in msgs/robot/tick as N grows (see X-25)",
     "S", 8, 1),
    ("X-14", "Scalability curve", "M5 + M6",
     "Tick compute time and throughput per robot-hour at 10, 25, 50 and 100 "
     "robots.",
     "Four data points recorded from real runs, not extrapolated",
     "S", 8, 1),
    ("X-22", "Trace hash strip", "M6 + M1",
     "Seed, git commit, config hash and rolling trace hash always visible in "
     "the status bar.",
     "Two runs with the same seed produce an identical rolling hash",
     "S", 7, 1),
    ("X-17", "Heterogeneous fleet", "M5 + M4",
     "Payload class, speed class and charge rate per robot; bids gated by "
     "capability feasibility.",
     "A robot that cannot carry a payload class never wins that task",
     "S", 7, 1),
    ("X-16", "Task priority classes and deadlines", "M5 + M4",
     "Priority classes with deadlines and pre-emption; SLA-miss rate as a KPI.",
     "SLA-miss KPI present and responsive to load",
     "S", 7, 1),
    ("X-18", "Battery feasibility veto", "M4",
     "A robot cannot win a task whose round trip plus reserve exceeds its "
     "state of charge.",
     "No robot ever strands; veto counter visible",
     "S", 6, 1),

    # --- Wave 2: M-effort, these four carry the demo
    ("X-12", "Counterfactual ghost fleet", "M5 + M1",
     "Same seed and identical task stream, a stop-and-wait baseline allocator "
     "rendered as ghosts on the SAME map, with a live divergence and task-lead "
     "counter.",
     "Live lead counter proves the >=20% completion-time reduction",
     "M", 10, 2),
    ("X-01", "Sovereign Agent Mode", "M4 + M2",
     "Each robot runs as a real OS process (multiprocessing) over a lossy "
     "message bus with no shared state. Built as a SWAPPABLE TRANSPORT behind "
     "the existing transport.py interface so the single-process path always "
     "keeps working.",
     "kill -9 on one agent process; fleet survives and recovers",
     "M", 10, 2),
    ("X-10", "Rogue-robot quarantine", "M4",
     "HMAC-signed CoordinationMessage, detection of bid inflation and of "
     "verdict non-compliance, automatic isolation of the offender.",
     "Rogue inflates bids, is detected and quarantined, throughput recovers",
     "M", 9, 2),
    ("X-20", "ML as congestion forecaster", "M3",
     "Predict cell occupancy at t+N from real recorded runs. Replaces the "
     "withdrawn action classifier that was trained on our own rule outputs.",
     "Forecast MAE reported against a held-out seed; advisory only",
     "M", 9, 2),

    # --- Wave 3: mechanism items contributed by external reviewers 3 and 4
    ("X-23", "Failure detector and gossip reservation GC", "M4",
     "10 Hz heartbeats. A miss of more than 200 ms causes adjacent peers to "
     "declare failure, gossip-invalidate the dead agent's active space-time "
     "reservations, and reroute around its last known coordinate.",
     "After kill -9 no ghost reservations remain; no deadlock follows",
     "S", 10, 3),
    ("X-24", "Deterministic lock arbitration", "M4",
     "Concurrent claims on the same cell tie-break on the immutable "
     "multi-attribute utility score U_i, then on unique agent sequence ID. "
     "Higher U takes the lock, lower yields.",
     "tests/test_race_condition.py - repeatable winner under simultaneity",
     "S", 8, 3),
    ("X-25", "Bounded radio range R_comm = 15 m", "M4",
     "Intent broadcast reaches only neighbours within 15 m, giving O(k) "
     "message complexity with k much smaller than N. Makes the Edge claim "
     "physically honest.",
     "msgs/robot/tick stays flat from 10 to 100 robots",
     "S", 9, 3),
    ("X-26", "Decision Inspector as the headline surface", "M1",
     "Pinned rail card showing the selected robot's intent, its space-time "
     "reservation, the per-term utility stack as a horizontal stacked bar, the "
     "verdict, and the NUMERIC winning margin over the runner-up bidder.",
     "Every number in the card traceable to a real coordination decision",
     "S", 9, 3),
    ("X-27", "Success-criteria instrumentation", "M6",
     "Two headline figures computed and displayed: collision count (target 0) "
     "and task-completion-time reduction versus the stop-and-wait baseline "
     "(target >= 20%).",
     "Both criteria met across >=10 seeds with 95% CI",
     "S", 10, 3),
]

DEFER = [
    ("X-05", "Congestion-aware time-expanded cost field", "M",
     "Correct but not load-bearing for the demo. AUTO-PROMOTES if X-20 lands "
     "early, because the forecaster needs a consumer."),
    ("X-06", "Deadlock avoidance via priority inheritance", "M",
     "Detection plus the graded ladder plus X-23 already give a recoverable "
     "system. Avoidance is a strengthening, not a gap."),
    ("X-08", "Regret-based SSI bidding with de-commitment", "M",
     "Strengthens N1 but the N1 claim is already restated honestly as "
     "incremental, so the exposure is closed without it."),
    ("X-11", "Operator authority with audit trail", "M",
     "Attractive for a defence sponsor but competes for the same build days as "
     "X-10, which is the stronger BEL claim."),
    ("X-15", "Sensing-reality knob", "M",
     "Noise and drift HOOKS are built into M5 now; the live UI controls are "
     "deferred. The sim-to-real question is answerable either way."),
    ("X-19", "UCB1 bandit over bid-weight vectors", "M",
     "Deferred, and claim N8 is DELETED rather than shipped soft. Nine airtight "
     "claims beat ten with one attackable one."),
    ("X-21", "Anytime planner with per-tick compute budget", "M",
     "The tick budget is MEASURED and displayed via X-04; the anytime fallback "
     "is not needed at these fleet sizes."),
]

REJECT = [
    ("X-07", "Bounded windowed CBS conflict-tree replan", "L",
     "All four reviewers ranked it last (impact/effort score 1.6) and the "
     "primary review explicitly says do not half-build it. We implement the "
     "ESCALATION HOOK and document the design, and answer the question "
     "verbally. A half-built optimal search is worse than a clearly scoped "
     "absence."),
    ("A-02", "RandomForest action classifier trained on rule outputs", "M",
     "DELETED OUTRIGHT. Behaviour-cloning our own rule engine caps the model at "
     "the rules and hands a judge the kill shot. Replaced by X-20."),
    ("N8", "Closed-loop analytics into bid weights, as originally specified", "M",
     "As written it was offline parameter tuning with a feedback arrow drawn on "
     "a slide. Claim withdrawn rather than defended."),
]

CUTS = [
    ("U-08", "Global search across all entity types",
     "Overlaps the command palette. A judge will never type a task ID."),
    ("U-05", "Command palette breadth",
     "Palette KEPT but trimmed to the 8 commands that actually appear in the "
     "demo script."),
    ("U-06", "Focus mode as a separate mode",
     "Selection-driven dimming gives 90% of the value for 10% of the work."),
    ("U-13", "Analytics screen chart count",
     "Cut to four charts: throughput, p95 latency, verdict distribution, "
     "conflicts per robot-minute."),
    ("S-06", "Scenario suite of six",
     "Ship three: rush_50, narrow_aisle_deadlock, blocked_aisle. "
     "battery_crisis and robot_failure become fault injections instead."),
    ("U-18", "Full ARIA and accessibility audit",
     "Keep focus rings and keyboard navigation; drop the audit. Not scored at "
     "SIH."),
    ("B-05", "GET /api/report and script-generated PDF deliverables",
     "Nothing on screen. Nobody opens a PDF during judging."),
    ("U-11", "Congestion heatmap as a decorative layer",
     "Kept ONLY if it drives a real cost field. Otherwise it is coloured "
     "noise."),
    ("A-02", "Action classifier (see REJECT table)",
     "Replaced by the X-20 congestion forecaster."),
    ("N8", "Closed-loop learning claim (see REJECT table)",
     "Withdrawn rather than shipped as tuning dressed as learning."),
]

NOVELTY = [
    ("N1", "Decentralised auction task allocation", "INCREMENTAL - stated honestly",
     "Contract Net Protocol (Smith, 1980); ST-SR-IA in the Gerkey and Mataric "
     "taxonomy. We name the prior art and compete on the measured result.",
     "Makespan vs single-shot auction and vs centralised greedy", 6),
    ("N2", "Intent broadcast with space-time reservation", "TABLE STAKES",
     "WHCA* (Silver, 2005) and Kiva-class commercial practice. Kept, but no "
     "longer claimed as novel.",
     "X-09 zero-violation counter makes it a measured invariant", 5),
    ("N3", "Deterministic replay", "GENUINE (engineering)",
     "Not a research contribution, but rare in multi-agent demos and it buys "
     "real auditability.",
     "X-22 rolling trace hash; two instances diffed live", 8),
    ("N4", "Deadlock detection with a graded response ladder", "INCREMENTAL",
     "Wait-for graph cycle detection is classical (Coffman conditions; Knapp "
     "1987). Zone-control deadlock handling is standard AGV literature (Kim "
     "and Tanchoco; Reveliotis).",
     "Graded ladder PROCEED/WAIT/YIELD/SLOW/REROUTE visible per decision", 6),
    ("N5", "Edge-AI advisory behind a hard safety firewall", "GENUINE",
     "This is runtime assurance, the Simplex architecture (Sha, 2001). Naming "
     "the prior art STRENGTHENS the claim because the contribution is the "
     "enforcement and its measurement, not the idea.",
     "Live veto counter; ML killed mid-run with no safety change; X-04", 9),
    ("N6", "Graceful degradation ladder FULL/NO_ML/NO_NETWORK", "INCREMENTAL",
     "Mode-based degradation is standard fault-tolerant design. It only counts "
     "if degradation is observed rather than declared.",
     "X-02 injected packet loss drives the transition", 7),
    ("N7", "Live digital-twin command centre", "DIFFERENTIATION, not novelty",
     "A UI is not a technical contribution. Claiming it as one invites a "
     "reviewer to discount the whole list. The defensible part is rendering "
     "the coordination REASONING.",
     "X-26 Decision Inspector with numeric winning margin", 9),
    ("N9", "Adversarial robot containment", "GENUINE - NEW",
     "Byzantine-tolerant multi-robot task allocation is an active research "
     "area and essentially absent from hackathon entries. For a defence PSU "
     "sponsor it is the most relevant claim we can make.",
     "X-10 rogue detected, quarantined, throughput recovers on stage", 9),
    ("N10", "Counterfactual co-simulation", "GENUINE - NEW",
     "Running the baseline and SWARMOS on an identical seed and task stream in "
     "the same frame, with live divergence, is an evidence technique no "
     "competing team will have built.",
     "X-12 ghost fleet plus X-13 statistical honesty", 10),
]

UIUX = [
    ("Everything reads mid-grey on near-black, so nothing is ranked",
     "Three-level ink system: primary #E8EDF2, secondary #9AA7B4, tertiary "
     "#63707D on surface #0F1419. Labels are ALWAYS tertiary, values ALWAYS "
     "primary."),
    ("Status colour used decoratively",
     "Saturated colour is reserved for STATE only: BLOCKED #E5484D, WAITING "
     "#F5A623, CHARGING #3E9BFF, MOVING neutral #C8D2DC, AVAILABLE #2F6F4E "
     "outline. Idle robots must be the quietest thing on the map."),
    ("50 robots become confetti",
     "Level of detail: below 0.6 zoom draw 4px dots with a heading tick only; "
     "above it draw footprint, heading wedge and ID. Never draw all 50 trails "
     "- only the selected robot and any robot in an active conflict."),
    ("Conflict indicators that never stop flashing",
     "One 400 ms pulse at conflict birth, then a static 2px ring until "
     "resolution. Continuous animation reads as a screensaver."),
    ("Six equal-weight KPI tiles",
     "One hero metric (tasks/min) at 40px tabular numerals, three supporting "
     "metrics at 20px, everything else behind the Analytics tab."),
    ("Empty states that just say 'No data'",
     "Every empty state names the cause and the next action, with the button "
     "inline: 'No incidents. Inject a blocked aisle from Simulation Lab to "
     "generate one.'"),
    ("Error states as red toasts",
     "Inline degraded banners bound to the mode ladder, naming the failing "
     "subsystem, the time since failure, and what still works."),
    ("Generic dark-dashboard layout",
     "Fixed three-zone operator layout: 68% map canvas, 22% right context rail "
     "(selection then incidents), 10% bottom command strip. No panel moves, "
     "ever. Operators memorise geometry."),
    ("Tooltips as the only explanation route",
     "Decision Inspector opens as a pinned rail card with the utility "
     "breakdown as a horizontal stacked bar and the winning margin called out "
     "numerically."),
    ("No visible eye entry point",
     "Top-left strip holds mode, seed, tick and the invariant counter. That is "
     "the first fixation and the only place carrying a coloured state chip "
     "when the system is healthy."),
]

SUCCESS = [
    ("Zero collisions", "0 space-time envelope overlaps, ever",
     "X-09 live counter plus an assertion in the benchmark harness across all "
     "scenarios and all seeds"),
    ("Task-completion-time reduction", ">= 20% versus a stop-and-wait baseline",
     "X-12 ghost fleet on an identical seed, X-13 paired per-seed deltas with "
     "95% CI over >= 10 seeds"),
    ("Real-time budget", "10 Hz tick sustained at 50 robots",
     "X-04 tick headroom panel; X-14 scalability curve to 100 robots"),
    ("Determinism", "Same seed produces an identical trace hash",
     "X-22 rolling hash in the status bar; two instances diffed live"),
    ("Communication scalability", "msgs/robot/tick flat as fleet size grows",
     "X-25 bounded radio range; X-03 plotted against fleet size"),
    ("Survives agent loss", "kill -9 an agent, fleet recovers, no deadlock",
     "X-01 real processes plus X-23 gossip reservation GC"),
]

CONVENTIONS = [
    ("Spatial model", "Continuous 2-D metres. M5 owns grid<->metre conversion"),
    ("Heading", "Radians in [0, 2*pi), 0 = +X axis"),
    ("Velocity", "Scalar speed in m/s, >= 0"),
    ("Battery", "Percentage 0-100"),
    ("Tick rate", "10 Hz, 100 ms budget per tick"),
    ("Robot IDs", "R### e.g. R001"),
    ("Task IDs", "T### e.g. T001"),
    ("Radio range", "R_comm = 15 m (X-25)"),
    ("Heartbeat", "10 Hz; failure declared after a 200 ms gap (X-23)"),
    ("Envelope", "schema_version '1.0'"),
    ("Interpreter",
     "/pkg/OSS-python-/3.12.0/x86_64-linux4.18-glibc2.28/bin/python3"),
]

LAWS = [
    "M5 simulation is the SINGLE authoritative source of robot state. Nothing "
    "else writes it.",
    "M4 coordination is the BINDING safety arbiter. Its verdict is final.",
    "M3 machine learning is ADVISORY ONLY and is never in the safety path.",
    "ONE canonical model per concept, defined in app/coordination/models.py.",
]

BUILD_ORDER = [
    ("1", "M5 simulation", "app/sim/",
     "Warehouse, robots, tasks, scenarios, seeded determinism, time controls, "
     "noise/drift hooks, X-14 scalability, X-17 heterogeneity, X-16 priorities"),
    ("2", "M6 database and analytics", "app/db/",
     "Persistence, trace recording, KPIs, X-13 multi-seed statistics, X-03 "
     "message accounting, X-22 trace hash, X-27 success criteria"),
    ("3", "M4 coordination Step 4B", "app/coordination/",
     "X-09 invariant monitor, X-18 battery veto, X-24 lock arbitration, X-25 "
     "bounded radio, X-02 impairment injection, X-23 failure detector and "
     "gossip GC, verdict ladder, then X-01 process transport and X-10 "
     "quarantine"),
    ("4", "M2 backend", "app/api/",
     "The 10 Hz tick loop, REST surface, /ws/fleet websocket"),
    ("5", "M1 frontend", "web/",
     "tokens.css and the fixed 68/22/10 shell FIRST, then Live Map, X-26 "
     "Decision Inspector, Simulation Lab, Analytics, then the 8-command "
     "palette"),
    ("6", "M3 machine learning", "app/ml/",
     "X-20 congestion forecaster trained on real recorded runs, X-04 compute "
     "budget measurement, firewall veto counter"),
    ("7", "X-12 ghost fleet", "app/sim/ + web/",
     "Stop-and-wait baseline co-simulated on the identical seed, ghost render "
     "layer, live divergence and task-lead counter"),
    ("8", "Verification and rehearsal", "tests/ + traces/",
     "Prove zero collisions and the >= 20% reduction across >= 10 seeds; "
     "rehearse the 5-minute demo including the kill -9 beat"),
]

DEMO = [
    ("0:00-0:20", "Opening frame",
     "Live Map already running with 50 robots, mode FULL, seed, tick and the "
     "safety counter at 0 in the top-left strip. No landing page, ever."),
    ("0:20-0:55", "Normal operation",
     "Start rush_50 from Simulation Lab. Click one robot to open the Decision "
     "Inspector and show the real utility stack and winning margin."),
    ("0:55-1:35", "Tension",
     "Inject blocked_aisle. Conflicts appear, the verdict ladder escalates, "
     "the invariant counter stays at 0."),
    ("1:35-2:20", "The hard failure",
     "kill -9 a real agent process on a visible terminal. Peers miss "
     "heartbeats past 200 ms, gossip-invalidate its reservations, reroute. "
     "Fleet degrades and recovers with zero violations."),
    ("2:20-3:00", "The adversary",
     "Enable the rogue robot. It inflates bids, is detected, is quarantined, "
     "throughput recovers on the chart."),
    ("3:00-3:40", "The reveal",
     "Turn on the ghost fleet. The stop-and-wait baseline runs on the same "
     "seed and visibly falls behind; the task-lead counter climbs."),
    ("3:40-4:20", "The proof",
     "Benchmarks screen: 10+ seeds, 95% CI error bars, the >= 20% paired "
     "delta, the flat msgs/robot/tick curve, the scalability curve to 100."),
    ("4:20-4:45", "Honesty",
     "Edge compute panel: model KB, inference p95, tick headroom. Kill the ML "
     "mid-run - safety is unchanged because ML is advisory."),
    ("4:45-5:00", "The close",
     "Replay the entire run from the trace hash and show an identical hash. "
     "Everything you just saw is reproducible from one seed."),
]


# --------------------------------------------------------------------------
# build
# --------------------------------------------------------------------------

def build(out_path):
    doc = Document()

    for s in doc.sections:
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)

    t = doc.add_heading("SWARMOS", level=0)
    for r in t.runs:
        r.font.color.rgb = ACCENT
    para(doc, "Final Implementation Feature List - the frozen build contract",
         size=13, bold=True)
    para(doc, "Edge-AI Based Distributed Fleet Coordination for Autonomous "
              "Mobile Robots in Smart Warehouses", size=11, italic=True)
    para(doc, "Smart India Hackathon 2026  |  Problem ID SIH26123  |  Bharat "
              "Electronics Limited  |  Theme: Robotics & Drones",
         size=10, color=MUTED)
    doc.add_paragraph()
    callout(doc, "This document is the authoritative build scope. It was frozen "
                 "after four independent external AI adversarial reviews were "
                 "ingested, cross-compared and triaged. Anything not listed "
                 "here is explicitly out of scope.", color=ACCENT)

    # ---------------- 1 how the scope was decided
    h1(doc, "1. How This Scope Was Decided")
    para(doc, "The project blueprint was sent to four independent AI reviewers "
              "with an adversarial four-voice prompt (national judge, "
              "multi-agent robotics professor, industrial-HMI designer, rival "
              "team lead). Their returns were cross-compared.")
    bullet(doc, "All four reviewers independently converged on the SAME set of "
                "22 additions. No fifth major idea appeared anywhere. The "
                "feature set is therefore treated as CLOSED.")
    bullet(doc, "Reviewers 3 and 4 contributed mechanism-level detail that 1 "
                "and 2 left abstract. Those five items were accepted and are "
                "listed as X-23 to X-27.")
    bullet(doc, "The central finding, accepted in full: three of the eight "
                "original novelty claims were soft and two were "
                "indefensible. The novelty register was rewritten.")
    bullet(doc, "Ten features were CUT to buy the build time for the "
                "additions. A shorter, sharper, fully-evidenced demo beats a "
                "broad one.")

    h2(doc, "1.1 Headline consequences")
    for x in [
        "The deliberate on-stage failure is a real kill -9 on a real OS "
        "process, not a simulated fault flag.",
        "The benchmark is no longer a bar chart. It is a live race against a "
        "stop-and-wait ghost baseline on the identical seed.",
        "The Edge-AI claim is carried by measured inference latency and a "
        "congestion forecaster. The action classifier trained on our own rule "
        "outputs is WITHDRAWN.",
        "The safety claim is carried by a zero-valued violation counter that "
        "is permanently on screen, not by an assertion buried in a test suite.",
        "The command centre is presented as DIFFERENTIATION, not as novelty.",
    ]:
        bullet(doc, x)

    # ---------------- 2 architectural laws
    h1(doc, "2. Architectural Laws (unchanged, non-negotiable)")
    for x in LAWS:
        numbered(doc, x)

    # ---------------- 3 conventions
    h1(doc, "3. Frozen Conventions")
    table(doc, ["Item", "Decision"], CONVENTIONS, widths=[1.8, 5.4], size=9)

    # ---------------- 4 success criteria
    h1(doc, "4. Success Criteria - what the build must PROVE")
    para(doc, "These are the published SIH26123 criteria plus the engineering "
              "targets that make them credible. Every one is measured on "
              "screen, not asserted.", italic=True)
    table(doc, ["Criterion", "Target", "How it is proven"], SUCCESS,
          widths=[1.7, 2.0, 3.5], size=9)

    # ---------------- 5 already done
    h1(doc, "5. Already Implemented and Unit-Tested (M4 Step 4A)")
    table(doc, ["ID", "Capability", "Module", "Test"], DONE,
          widths=[0.6, 3.6, 0.8, 2.2], size=9)

    # ---------------- 6 the build list
    h1(doc, "6. THE BUILD LIST - Features To Be Implemented")
    para(doc, "%d accepted additions. Effort: S = under a day, M = 1 to 3 "
              "days. Impact is judge impact on a 1 to 10 scale. Wave 1 is "
              "cheap evidence, Wave 2 carries the demo, Wave 3 are the "
              "mechanism items that make Waves 1 and 2 actually work."
         % len(BUILD), italic=True)

    for wave, title, note in [
        (1, "Wave 1 - S-effort evidence (about one week)",
         "Every item here converts a weak claim into a number a judge can "
         "watch."),
        (2, "Wave 2 - M-effort demo carriers (about two weeks)",
         "These four are the demo. If time runs short, everything else goes "
         "before these do."),
        (3, "Wave 3 - mechanism items from reviewers 3 and 4",
         "Small, cheap, and each one closes a real hole rather than adding "
         "surface."),
    ]:
        h2(doc, title)
        para(doc, note, size=9.5, italic=True, color=MUTED)
        rows = [(i, n, m, w, a, e, imp)
                for (i, n, m, w, a, e, imp, wv) in BUILD if wv == wave]
        table(doc, ["ID", "Feature", "Module", "What gets built",
                    "Acceptance test", "Eff", "Imp"],
              rows, widths=[0.5, 1.2, 0.7, 2.4, 1.7, 0.35, 0.35], size=8)

    total_s = sum(1 for b in BUILD if b[5] == "S")
    total_m = sum(1 for b in BUILD if b[5] == "M")
    callout(doc, "Build total: %d features - %d S-effort, %d M-effort. "
                 "Plus the 6 items already complete." % (len(BUILD), total_s,
                                                         total_m),
            color=GOOD)

    # ---------------- 7 novelty register
    h1(doc, "7. Final Novelty Register - 9 Claims, All Defensible")
    para(doc, "The original eight claims were re-adjudicated against prior "
              "art. N8 was DELETED. N9 and N10 are new. Where a claim is "
              "standard practice we now say so and name it - that is what "
              "makes the remaining claims survive a professor.", italic=True)
    table(doc, ["ID", "Claim", "Verdict", "Prior art / reasoning",
                "Evidence on screen", "Imp"],
          NOVELTY, widths=[0.4, 1.4, 1.1, 2.3, 1.7, 0.3], size=8)
    callout(doc, "We claim NINE things and can prove all nine. We do not claim "
                 "a tenth.", color=GOOD)

    # ---------------- 8 deferred
    h1(doc, "8. Deferred - Correct, But Not In This Build")
    para(doc, "Each of these is a genuine improvement. None is load-bearing "
              "for the success criteria or the demo. They are the stretch "
              "list, attempted only after the build list is complete, tested "
              "and rehearsed.", italic=True)
    table(doc, ["ID", "Item", "Eff", "Why deferred"], DEFER,
          widths=[0.5, 2.2, 0.35, 4.1], size=9)

    # ---------------- 9 rejected
    h1(doc, "9. Rejected - With Reasons On The Record")
    para(doc, "Written down so that the rejection is defensible if a judge "
              "asks the same question. Being able to say WHY something was not "
              "built is stronger than pretending it was never considered.",
         italic=True)
    table(doc, ["ID", "Item", "Eff", "Why rejected"], REJECT,
          widths=[0.5, 2.2, 0.35, 4.1], size=9)

    # ---------------- 10 cuts
    h1(doc, "10. Cut From The Original Inventory")
    para(doc, "These cuts fund the additions. The marginal mark was lower than "
              "the marginal hour.", italic=True)
    table(doc, ["ID", "Cut or descoped", "Reason"], CUTS,
          widths=[0.5, 2.3, 4.4], size=9)

    # ---------------- 11 uiux
    h1(doc, "11. UI/UX Corrections - Anti Vibe-Coded Rules")
    para(doc, "Each row is a concrete visual defect and its specific remedy. "
              "All token values live in web/styles/tokens.css. No ad-hoc hex "
              "value may appear anywhere else in the codebase.", italic=True)
    table(doc, ["Defect", "Specific fix"], UIUX, widths=[2.4, 4.8], size=9)

    h2(doc, "11.1 Standing rules")
    for x in [
        "NO FAKE DATA, EVER. Every number on screen comes from a real run.",
        "Every control performs a real action. No decorative buttons.",
        "Purposeful motion only, 120-200 ms. Nothing loops.",
        "One UI font plus one tabular-numeric monospace for all figures.",
        "No emoji, no gradient soup, no purple-on-black.",
        "Designed empty, loading, error and degraded states for every panel.",
        "Keyboard-first: Ctrl/Cmd+K palette (8 commands), arrow keys cycle "
        "robots, space pauses.",
        "System health is reported honestly, including when it is bad.",
    ]:
        bullet(doc, x)

    # ---------------- 12 build order
    h1(doc, "12. Build Order")
    para(doc, "The system is kept runnable at the end of every step rather "
              "than integrated once at the end.", italic=True)
    table(doc, ["#", "Step", "Path", "Contents"], BUILD_ORDER,
          widths=[0.3, 1.5, 1.2, 4.2], size=9)

    h2(doc, "12.1 The one hedge")
    callout(doc, "X-01 Sovereign Agent Mode is the highest-risk item. It is "
                 "built as a SWAPPABLE TRANSPORT behind the existing "
                 "app/coordination/transport.py interface, with the "
                 "single-process tick loop working first. If process mode runs "
                 "long we still have a complete working demo and lose exactly "
                 "one claim instead of the whole system.")

    # ---------------- 13 demo
    h1(doc, "13. Demo Script - What The Build Is Aimed At")
    para(doc, "Five minutes. Evidence, tension, recovery, comparison. Every "
              "beat exists because a feature in section 6 makes it possible.",
         italic=True)
    table(doc, ["Time", "Beat", "Screen and action"], DEMO,
          widths=[0.9, 1.3, 5.0], size=9)

    # ---------------- 14 repo layout
    h1(doc, "14. Repository Layout")
    mono(doc,
         "app/coordination/   M4  canonical models, messages, transport,\n"
         "                        registries, reservation, conflict, auction,\n"
         "                        invariant monitor, quarantine, failure det.\n"
         "app/sim/            M5  warehouse, robots, tasks, scenarios, clock,\n"
         "                        baseline stop-and-wait allocator, ghost co-sim\n"
         "app/db/             M6  sqlite persistence, KPIs, trace recorder,\n"
         "                        benchmark harness, statistics\n"
         "app/ml/             M3  congestion forecaster, compute budget meter\n"
         "app/api/            M2  FastAPI app, 10 Hz tick loop, /ws/fleet\n"
         "web/                M1  index.html, styles/tokens.css, js/ modules\n"
         "tests/                  pytest suite\n"
         "docs/                   generators and design records\n"
         "traces/                 recorded runs for replay\n"
         "run.sh                  single entry point\n")

    doc.add_paragraph()
    para(doc, "End of implementation list. Generated from "
              "docs/gen_implementation_list_docx.py - regenerate after any "
              "scope change.", size=9, italic=True, color=MUTED)

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    out = os.path.join(os.path.dirname(here),
                       "SWARMOS_Final_Implementation_Feature_List.docx")
    path = build(out)
    d = Document(path)
    print("Wrote: %s" % path)
    print("Paragraphs: %d   Tables: %d" % (len(d.paragraphs), len(d.tables)))
    print("Build list: %d   Done: %d   Deferred: %d   Rejected: %d   Cuts: %d"
          % (len(BUILD), len(DONE), len(DEFER), len(REJECT), len(CUTS)))
    print("Novelty claims: %d   UI/UX fixes: %d" % (len(NOVELTY), len(UIUX)))

# File contains AI-generated response based on internal company sources
